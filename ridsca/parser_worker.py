"""Private parser entrypoint. No incoming filenames become commands or output paths."""

import hashlib
import json
from pathlib import Path
import subprocess
import sys
from .config import MAX_TEXT
from .parsing import AUDIO, VIDEO, IMAGES, clamp_text
from .schema import ParsedInput, Check


def check(result, component, status, detail):
    result.checks.append(Check(component=component, status=status, detail=detail, source=result.source))


def parse(path, source):
    ext = path.suffix.lower()
    r = ParsedInput(source=source, kind="document")
    if ext in {".txt", ".md", ".json", ".html", ".htm"}:
        text = path.read_text(encoding="utf-8-sig")
        if "\x00" in text:
            raise ValueError("Binary content is not accepted as text")
        if ext == ".json":
            text = json.dumps(json.loads(text), ensure_ascii=False)
        elif ext in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(text, "html.parser")
            for tag in soup(["script", "style", "iframe", "object"]):
                tag.decompose()
            links = [a.get("href", "") for a in soup.find_all("a")][:100]
            text = soup.get_text(" ", strip=True) + "\n" + "\n".join(links)
        r.text = text
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        parts = []
        for i, page in enumerate(reader.pages[:20]):
            text = page.extract_text() or ""
            if not text.strip():
                check(
                    r,
                    "PDF extraction",
                    "partial",
                    f"Page {i + 1} has no extractable text; scan OCR was not performed.",
                )
            parts.append(f"[Page {i + 1}]\n{text}")
        r.text = "\n".join(parts)
        if len(reader.pages) > 20:
            check(r, "PDF extraction", "partial", "Only the first 20 pages were read.")
    elif ext == ".docx":
        from docx import Document

        doc = Document(path)
        r.text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        )
        check(
            r,
            "DOCX coverage",
            "partial",
            "Body paragraphs and tables only; embedded images, headers and comments not read.",
        )
    elif ext == ".pptx":
        from pptx import Presentation

        slides = list(Presentation(path).slides)
        r.text = "\n".join(
            f"[Slide {i + 1}] " + " ".join(s.text for s in slide.shapes if s.has_text_frame)
            for i, slide in enumerate(slides[:30])
        )
        check(
            r,
            "PPTX coverage",
            "partial",
            "Text boxes on up to 30 slides only; embedded media and notes not read.",
        )
    elif ext in IMAGES:
        import cv2
        import numpy as np
        from PIL import Image
        import pytesseract

        r.kind = "image"
        Image.MAX_IMAGE_PIXELS = 16_000_000
        with Image.open(path) as im:
            rgb = im.convert("RGB")
            rgb.thumbnail((1800, 1800))
            try:
                r.text = pytesseract.image_to_string(rgb, lang="eng+hin", timeout=15)
                check(r, "OCR", "complete", "English/Hindi OCR; transcription errors are possible.")
            except Exception:
                check(r, "OCR", "unavailable", "Tesseract or language data unavailable; no text inferred.")
            detector = cv2.QRCodeDetector()
            arr = cv2.cvtColor(np.asarray(rgb), cv2.COLOR_RGB2BGR)
            ok, payloads, _, _ = detector.detectAndDecodeMulti(arr)
            if ok:
                r.qr = [s[:2048] for s in payloads if s][:10]
            else:
                payload, _, _ = detector.detectAndDecode(arr)
                r.qr = [payload[:2048]] if payload else []
            check(
                r, "QR decoding", "complete", f"Decoded {len(r.qr)} payload(s); never opened automatically."
            )
    elif ext in AUDIO | VIDEO:
        r.kind = "audio" if ext in AUDIO else "video"
        proc = subprocess.run(
            [
                "ffprobe",
                "-v",
                "error",
                "-protocol_whitelist",
                "file,pipe",
                "-show_entries",
                "format=duration:stream=codec_type,width,height",
                "-of",
                "json",
                str(path),
            ],
            capture_output=True,
            timeout=10,
            check=True,
        )
        meta = json.loads(proc.stdout)
        streams = meta.get("streams", [])
        duration = float(meta.get("format", {}).get("duration", 0))
        if duration <= 0 or duration > 3600:
            raise ValueError("Invalid duration or media longer than one hour")
        for s in streams:
            if s.get("width", 0) * s.get("height", 0) > 8_300_000:
                raise ValueError("Video exceeds supported resolution")
        r.metadata = {"duration": duration, "has_audio": any(s.get("codec_type") == "audio" for s in streams)}
        check(r, "Media container", "complete", f"Container decoded; duration {duration:.1f}s.")
    elif ext == ".apk":
        r.kind = "apk"
        r.metadata["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        try:
            proc = subprocess.run(
                ["aapt", "dump", "badging", str(path)], capture_output=True, timeout=15, check=True
            )
            import re

            output = proc.stdout.decode("utf-8", errors="replace")[:100000]
            package = re.search(r"package: name='([^']+)'", output)
            version = re.search(r"versionName='([^']+)'", output)
            permissions = re.findall(r"uses-permission(?:-sdk-\d+)?: name='([^']+)'", output)
            r.metadata.update(
                package=package.group(1) if package else "unknown",
                version=version.group(1) if version else "unknown",
                permissions=permissions[:200],
            )
            check(r, "APK metadata", "complete", "Manifest inspected with aapt; no APK code executed.")
        except Exception:
            check(
                r,
                "APK metadata",
                "unavailable",
                "aapt missing or unable to decode manifest; SHA-256 is still available.",
            )
    r.text, truncated = clamp_text(r.text)
    if truncated:
        check(r, "Text coverage", "partial", f"Extracted text limited to {MAX_TEXT} characters.")
    return r


def main():
    try:
        # Linux defense in depth. This is resource isolation, not a kernel sandbox.
        try:
            import resource

            resource.setrlimit(resource.RLIMIT_CPU, (40, 40))
            resource.setrlimit(resource.RLIMIT_FSIZE, (32 * 1024 * 1024, 32 * 1024 * 1024))
            resource.setrlimit(resource.RLIMIT_NOFILE, (128, 128))
        except (ImportError, ValueError):
            pass
        result = parse(Path(sys.argv[1]), sys.argv[2]).model_dump()
    except Exception as exc:
        # Do not leak paths, document text, or subprocess diagnostics.
        result = {"error": f"Unable to parse this file ({type(exc).__name__})."}
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
