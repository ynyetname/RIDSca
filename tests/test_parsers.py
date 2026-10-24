import io
import shutil
import wave
import zipfile
import pytest
from ridsca.parsing import parse_upload
from ridsca.config import Settings
from ridsca.service import Analyzer


def test_pdf_text_extraction():
    # A minimal PDF with a valid page/content stream, independent of reportlab.
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 30 250 Td (Send your OTP) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    parsed = parse_upload("a.pdf", out.getvalue(), "a")
    assert "Send your OTP" in parsed.text


def test_docx_paragraph_and_table():
    from docx import Document

    d = Document()
    d.add_paragraph("hello")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Send your OTP"
    b = io.BytesIO()
    d.save(b)
    assert "Send your OTP" in parse_upload("a.docx", b.getvalue(), "a").text


def test_pptx_text():
    from pptx import Presentation

    p = Presentation()
    p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "Send your OTP"
    b = io.BytesIO()
    p.save(b)
    assert "Send your OTP" in parse_upload("a.pptx", b.getvalue(), "a").text


def test_qr_payload_routes_to_url_check():
    import cv2
    import numpy as np

    qr = cv2.QRCodeEncoder_create().encode("http://127.0.0.1/private")
    qr = np.pad(qr, 4, constant_values=255)
    qr = cv2.resize(qr, None, fx=10, fy=10, interpolation=cv2.INTER_NEAREST)
    data = cv2.imencode(".png", qr)[1].tobytes()
    report = Analyzer(Settings()).analyze(uploads=[("qr.png", data)], consent=True)
    assert any(c.component == "QR decoding" for c in report.checks)
    assert any(c.component == "URL network" and c.status == "skipped" for c in report.checks)


@pytest.mark.skipif(not shutil.which("ffprobe"), reason="ffprobe not installed")
def test_wav_container():
    b = io.BytesIO()
    with wave.open(b, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(16000)
        w.writeframes(b"\x00\x00" * 16000)
    p = parse_upload("voice.wav", b.getvalue(), "a")
    assert p.kind == "audio"
    assert p.metadata["has_audio"]


def test_apk_hash_without_execution():
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("AndroidManifest.xml", '<manifest package="example.synthetic"/>')
    p = parse_upload("test.apk", b.getvalue(), "a")
    assert p.kind == "apk" and len(p.metadata["sha256"]) == 64
